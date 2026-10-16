"""Fill a workspace with demo documents for the weekly-sales-report scenario.

    python misc/make_demo_data.py [workspace_dir]

Creates an example report (docx) and sales exports (xlsx) dated in the previous,
current and an older week, so "last week's files" is a meaningful filter."""

from __future__ import annotations

import os
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook


CLIENTS = ["ООО Ромашка", "АО Вектор", "ИП Смирнов", "ООО Северсталь-Сервис", "ООО Глобус", "АО Техноресурс"]


def main() -> None:
    root = Path(sys.argv[1] if len(sys.argv) > 1 else "workspace").resolve()
    today = date.today()
    this_week = today - timedelta(days=today.weekday())
    prev_week = this_week - timedelta(days=7)

    example = root / "Отчёты" / "Пример" / "Отчёт_пример.docx"
    example.parent.mkdir(parents=True, exist_ok=True)
    _example_report(example, prev_week - timedelta(days=7))

    sales = root / "Отчёты" / "Продажи"
    sales.mkdir(parents=True, exist_ok=True)

    rng = random.Random(2026)
    for day in (prev_week - timedelta(days=5), prev_week + timedelta(days=1), prev_week + timedelta(days=3), this_week + timedelta(days=1)):
        if day > today:
            continue

        path = sales / f"Продажи_{day.isoformat()}.xlsx"
        _sales_export(path, day, rng)

        # fs.list filters by modification time; make it match the export's day.
        stamp = datetime.combine(day, datetime.min.time()).replace(hour=18).timestamp()
        os.utime(path, (stamp, stamp))

    inbox = root / "Входящие"
    inbox.mkdir(parents=True, exist_ok=True)
    (inbox / "Счёт_45_Ромашка.txt").write_text("Счёт на оплату № 45 от 06.10.2026. Поставщик ООО Ромашка. Сумма 120 000 руб.", encoding="utf-8")
    (inbox / "Договор_поставки_Вектор.txt").write_text("Договор поставки № 12/26. АО Вектор (Поставщик) и ООО Глобус (Покупатель) заключили настоящий договор…", encoding="utf-8")
    (inbox / "Акт_сверки_Смирнов.txt").write_text("Акт сверки взаимных расчётов между ИП Смирнов и ООО Глобус за III квартал 2026 года.", encoding="utf-8")

    print(f"demo data in {root}")


def _example_report(path: Path, week_start: date) -> None:
    document = Document()
    document.styles["Normal"].font.size = Pt(11)

    document.add_heading("Еженедельный отчёт по продажам", level=1)
    week_end = week_start + timedelta(days=6)
    document.add_paragraph(f"Период: {week_start:%d.%m.%Y} – {week_end:%d.%m.%Y}")
    document.add_paragraph("Выручка за неделю: 1 245 000 руб.")
    document.add_paragraph("Количество заказов: 87")

    document.add_paragraph("Крупнейшие клиенты недели:")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Клиент"
    table.rows[0].cells[1].text = "Сумма, руб."

    for client, amount in (("ООО Ромашка", "410 000"), ("АО Вектор", "305 000"), ("ИП Смирнов", "198 000")):
        row = table.add_row()
        row.cells[0].text = client
        row.cells[1].text = amount

    document.add_paragraph("Подготовил: отдел продаж")
    document.save(str(path))


def _sales_export(path: Path, day: date, rng: random.Random) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Продажи"
    sheet.append(["Дата", "Номер заказа", "Клиент", "Сумма, руб."])

    for i in range(rng.randint(8, 14)):
        sheet.append([day.isoformat(), f"З-{day:%m%d}-{i + 1:03d}", rng.choice(CLIENTS), rng.randrange(5_000, 120_000, 500)])

    workbook.save(str(path))


if __name__ == "__main__":
    main()
