"""Turn office files into plain text/markdown an LLM can read."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from common.tools import ToolError


TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log", ".xml", ".html", ".htm"}
MAX_SHEET_ROWS = 500


def read_document(path: Path) -> str:
    if not path.is_file():
        raise ToolError(f"Файл `{path.name}` не найден")

    suffix = path.suffix.lower()

    if suffix == ".docx":
        return _read_docx(path)

    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)

    if suffix == ".pdf":
        return _read_pdf(path)

    if suffix in TEXT_SUFFIXES:
        return _read_text(path)

    raise ToolError(f"Формат `{suffix}` пока не поддерживается")


def cell_to_text(value) -> str:
    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return str(value).replace("\n", " ").strip()


def markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""

    width = max(len(r) for r in rows)
    padded = [r + [""] * (width - len(r)) for r in rows]

    lines = ["| " + " | ".join(padded[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in padded[1:]]

    return "\n".join(lines)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []

    # Walk the body in order so tables stay where they are in the text.
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]

        if tag == "p":
            text = "".join(node.text or "" for node in block.iter() if node.tag.endswith("}t"))
            if text.strip():
                parts.append(text)

        elif tag == "tbl":
            rows = []
            for row in block.iter():
                if not row.tag.endswith("}tr"):
                    continue

                cells = []
                for cell in row:
                    if cell.tag.endswith("}tc"):
                        cells.append(" ".join(n.text or "" for n in cell.iter() if n.tag.endswith("}t")).strip())

                rows.append(cells)

            parts.append(markdown_table(rows))

    return "\n\n".join(parts)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []

    for sheet in workbook.worksheets:
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                rows.append([f"… ещё строки (показаны первые {MAX_SHEET_ROWS})"])
                break

            cells = [cell_to_text(v) for v in row]
            if any(cells):
                rows.append(cells)

        parts.append(f"## Лист «{sheet.title}»\n\n{markdown_table(rows)}")

    workbook.close()

    return "\n\n".join(parts)


def _read_pdf(path: Path) -> str:
    import pdfplumber

    with pdfplumber.open(str(path)) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]

    text = "\n\n".join(p for p in pages if p.strip())
    if not text:
        raise ToolError("В PDF нет текстового слоя (похоже на скан)")

    return text


def _read_text(path: Path) -> str:
    raw = path.read_bytes()

    # Office exports on Russian Windows are often cp1251.
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)

        except UnicodeDecodeError:
            continue

    return raw.decode("utf-8", errors="replace")
