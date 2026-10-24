"""Turn a filled-in example report into a reusable docxtpl template.

What varies (a date, a total, the rows of a table) is decided by a model outside this
tool: the agent in chat, a step of its own in a task. The tool checks that markup
against the document and edits the document in code, so formatting and layout stay
intact."""

from __future__ import annotations

import asyncio
import copy
import re
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

from common.tools import ToolContext, ToolError, ToolInput, ToolOutput, tool
from common.tools.fields import WorkspacePath


_NAME_RE = re.compile(r"^[a-z][a-z0-9_]{0,40}$")
_PLACEHOLDER = re.compile(r"\{\{ [a-z][a-z0-9_]* \}\}")


class TemplateField(BaseModel):
    name: str = Field(description="Имя поля латиницей snake_case")
    description: str = Field(description="Что это за значение")
    example: str = Field(description="ТОЧНЫЙ фрагмент текста документа, который станет полем")


class TemplateTable(BaseModel):
    table: int = Field(ge=1, description="Номер таблицы: «Таблица T2» в тексте документа → 2")
    name: str = Field(description="Имя таблицы латиницей snake_case")
    header_rows: int = Field(ge=0, description="Сколько строк шапки остаются как есть")
    columns: list[str] = Field(min_length=1, description="Ключи колонок по порядку, латиницей snake_case")


class Input(ToolInput):
    example: WorkspacePath = Field(description="Готовый отчёт-пример .docx")
    name: str = Field(description="Имя шаблона латиницей, например weekly_sales")
    fields: list[TemplateField] = Field(default_factory=list, description="Значения, которые меняются от раза к разу")
    tables: list[TemplateTable] = Field(default_factory=list, description="Таблицы, строки которых будут другими")


class TableOut(BaseModel):
    name: str
    columns: list[str]


class Output(ToolOutput):
    template: str
    fields: list[TemplateField]
    tables: list[TableOut]
    example_context: dict[str, Any]


@tool(
    name="docs.make_template",
    title="Сделать шаблон из примера",
    card="Сделать шаблон «{name}» из примера «{example}»",
    description=(
        "Create a reusable .docx template from an example report by the given markup: "
        "`fields` (exact text fragments that vary) become {{ fields }}, `tables` (numbered "
        "as «Таблица T1» in docs.read) become row loops. Returns the template path, its "
        "fields and the example's values as a ready context for docs.fill_template."
    ),
)
async def make_template(inp: Input, ctx: ToolContext) -> Output:
    if not _NAME_RE.match(inp.name):
        raise ToolError("Имя шаблона: латиница в нижнем регистре, цифры и _, например weekly_sales")

    source = ctx.read_path(inp.example)
    if source.suffix.lower() != ".docx" or not source.is_file():
        raise ToolError(f"Нужен существующий .docx, получено `{inp.example}`")

    if not inp.fields and not inp.tables:
        raise ToolError("Разметка пуста: укажите fields или tables — что в примере меняется от раза к разу")

    # A dry run leaves data:templates alone: the template goes to the run's own outputs,
    # where it can be looked at (and filled by the next step) all the same.
    if ctx.dry_run:
        target = ctx.output_path(f"{inp.name}.docx")
        ctx.log("Пробный запуск: шаблон не сохранён в data:templates, он в результатах запуска")
    else:
        target = ctx.write_path(f"data:templates/{inp.name}.docx")

    result = await asyncio.to_thread(_make, source, target, inp.fields, inp.tables)
    result["template"] = ctx.to_ref(target)

    return Output.model_validate(result)


def _make(source: Path, target: Path, fields: list[TemplateField], tables: list[TemplateTable]) -> dict[str, Any]:
    from docx import Document

    document = Document(str(source))

    problems = _check(document, fields, tables)
    if problems:
        # Every problem at once: the caller fixes the markup in one go.
        raise ToolError("Разметка не подходит к документу:\n- " + "\n- ".join(problems))

    target.parent.mkdir(parents=True, exist_ok=True)

    return _apply_markup(document, target, fields, tables)


def _check(document, fields: list[TemplateField], tables: list[TemplateTable]) -> list[str]:
    problems: list[str] = []
    names: set[str] = set()

    def name_ok(name: str, what: str) -> None:
        if not _NAME_RE.match(name):
            problems.append(f"{what} «{name}»: имя латиницей snake_case, например report_date")
        elif name in names:
            problems.append(f"{what} «{name}»: имя уже занято")
        names.add(name)

    loops = {spec.table - 1 for spec in tables}
    count = len(document.tables)

    for spec in tables:
        name_ok(spec.name, "Таблица")
        if spec.table > count:
            problems.append(f"Таблицы T{spec.table} нет: в документе таблиц — {count}")
            continue

        rows = document.tables[spec.table - 1].rows
        if spec.header_rows >= len(rows):
            problems.append(f"T{spec.table}: строк {len(rows)}, после шапки из {spec.header_rows} не остаётся строки данных")
            continue

        width = len(rows[spec.header_rows].cells)
        if len(spec.columns) > width:
            problems.append(f"T{spec.table}: колонок {width}, а ключей {len(spec.columns)}")

        for column in spec.columns:
            if not _NAME_RE.match(column):
                problems.append(f"T{spec.table}: ключ колонки «{column}» — латиницей snake_case")

    texts = [p.text for p in _text_paragraphs(document, loops)]
    for field in fields:
        name_ok(field.name, "Поле")
        example = field.example.strip()
        if not example:
            problems.append(f"Поле «{field.name}»: пустой example")
        elif not any(example in text for text in texts):
            problems.append(f"Поле «{field.name}»: фрагмента «{example}» нет в тексте документа (вне размеченных таблиц)")

    return problems


def _text_paragraphs(document, skip_tables: set[int]) -> list:
    """Body paragraphs and those in tables that stay as they are."""
    paragraphs = list(document.paragraphs)
    for index, table in enumerate(document.tables):
        if index in skip_tables:
            continue

        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)

    return paragraphs


def _apply_markup(document, target: Path, fields: list[TemplateField], tables: list[TemplateTable]) -> dict[str, Any]:
    example_context: dict[str, Any] = {}

    for spec in tables:
        rows = document.tables[spec.table - 1].rows
        example_context[spec.name] = [
            {key: (row.cells[i].text.strip() if i < len(row.cells) else "") for i, key in enumerate(spec.columns)}
            for row in rows[spec.header_rows:]
        ]
        _make_row_loop(document.tables[spec.table - 1], spec.header_rows, spec.name, spec.columns)

    paragraphs = _text_paragraphs(document, {spec.table - 1 for spec in tables})
    missing: list[str] = []

    # The longest fragment first: «01.01.2025» becomes a field before «2025» could cut it.
    for field in sorted(fields, key=lambda f: len(f.example.strip()), reverse=True):
        example = field.example.strip()
        placeholder = "{{ " + field.name + " }}"
        replaced = [_replace_in_runs(paragraph, example, placeholder) for paragraph in paragraphs]

        if not any(replaced):
            missing.append(
                f"Поле «{field.name}»: фрагмент «{example}» остался только внутри другого поля "
                "или в тексте ссылки — возьми фрагмент, который стоит в тексте отдельно"
            )

        example_context[field.name] = example

    if missing:
        raise ToolError("Разметка не подходит к документу:\n- " + "\n- ".join(missing))

    document.save(str(target))

    return {
        "fields": [field.model_dump() for field in fields],
        "tables": [{"name": spec.name, "columns": spec.columns} for spec in tables],
        "example_context": example_context,
    }


def _make_row_loop(table, header: int, name: str, columns: list[str]) -> None:
    """Keep one data row as the pattern and wrap it in docxtpl `{%tr for %}` rows."""
    rows = table.rows
    pattern = rows[header]

    for i, cell in enumerate(pattern.cells):
        cell.text = "{{ row." + columns[i] + " }}" if i < len(columns) else ""

    for row in list(rows[header + 1:]):
        row._tr.getparent().remove(row._tr)

    open_tr = copy.deepcopy(pattern._tr)
    close_tr = copy.deepcopy(pattern._tr)
    pattern._tr.addprevious(open_tr)
    pattern._tr.addnext(close_tr)

    _set_row_tag(open_tr, "{%tr for row in " + name + " %}")
    _set_row_tag(close_tr, "{%tr endfor %}")


def _set_row_tag(tr, tag: str) -> None:
    texts = [node for node in tr.iter() if node.tag.endswith("}t")]
    for i, node in enumerate(texts):
        node.text = tag if i == 0 else ""


def _replace_in_runs(paragraph, example: str, placeholder: str) -> bool:
    """Replace `example` where it stands, even split across runs: the placeholder takes
    the first of them, the rest of the paragraph keeps its runs and their formatting.
    Text of links and fields (not runs of the paragraph) is left alone."""
    runs = paragraph.runs
    joined = "".join(run.text for run in runs)
    # Inside a placeholder made a moment ago («1» in «{{ day_1 }}») is not the document's text.
    taken = [match.span() for match in _PLACEHOLDER.finditer(joined)]
    starts = [
        match.start() for match in re.finditer(re.escape(example), joined)
        if not any(a < match.end() and match.start() < b for a, b in taken)
    ]

    # From the end: what is replaced later does not move what comes before.
    for start in reversed(starts):
        end, position, first = start + len(example), 0, True

        for run in runs:
            run_start, run_end = position, position + len(run.text)
            position = run_end
            if run_end <= start or run_start >= end:
                continue

            text = run.text
            cut_from, cut_to = max(start, run_start) - run_start, min(end, run_end) - run_start
            run.text = text[:cut_from] + (placeholder if first else "") + text[cut_to:]
            first = False

    return bool(starts)
