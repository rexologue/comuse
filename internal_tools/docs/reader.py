"""Office files as text: what docs.read gives a model step or the agent. Tables become
markdown, numbered «Таблица T1», «Таблица T2»… — the numbers docs.make_template takes."""

from __future__ import annotations

import itertools
from datetime import date, datetime
from pathlib import Path

from common.tools import ToolError


TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log", ".xml", ".html", ".htm"}
MAX_SHEET_ROWS = 500


def read_document(path: Path) -> str:
    if not path.is_file():
        raise ToolError(f"Файл `{path.name}` не найден")

    suffix = path.suffix.lower()

    if suffix == ".docx":
        return _read_docx(path)

    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)

    if suffix == ".pdf":
        return _read_pdf(path)

    if suffix in TEXT_SUFFIXES:
        return _read_text(path)

    if suffix == ".eml":
        return _read_eml(path)

    raise ToolError(f"Формат `{suffix}` пока не поддерживается")


def cell_to_text(value) -> str:
    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return str(value).replace("\n", " ").strip()


def markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""

    width = max(len(r) for r in rows)
    padded = [r + [""] * (width - len(r)) for r in rows]

    lines = ["| " + " | ".join(padded[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in padded[1:]]

    return "\n".join(lines)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    # Tables are numbered as python-docx counts them (body level), so «T2» means
    # document.tables[1] to docs.make_template.
    tables = itertools.count(1)
    parts = _blocks(document.element.body, lambda: f"Таблица T{next(tables)}")

    # Headers and footers hold what a body does not (the company, the signature line, a
    # table of requisites): once each, however many sections and kinds of page repeat them.
    def said(parts_of) -> list[str]:
        seen, found = set(), []
        for section in document.sections:
            for part in parts_of(section):
                if part.is_linked_to_previous:
                    continue
                text = "\n".join(_blocks(part._element, lambda: "Таблица"))
                if text and text not in seen:
                    seen.add(text)
                    found.append(text)
        return found

    headers = said(lambda section: (section.first_page_header, section.header, section.even_page_header))
    footers = said(lambda section: (section.first_page_footer, section.footer, section.even_page_footer))

    return "\n\n".join(
        [f"Верхний колонтитул: {text}" for text in headers] + parts + [f"Нижний колонтитул: {text}" for text in footers]
    )


def _blocks(parent, table_title) -> list[str]:
    """The paragraphs and tables under `parent` in their order; a form's fields (content
    controls, w:sdt) are read as the text they hold."""
    parts: list[str] = []

    for block in parent.iterchildren():
        tag = _tag(block)

        if tag == "p":
            text = _paragraph_text(block)
            if text.strip():
                parts.append(text)

        elif tag == "tbl":
            rows = [
                [" ".join(t for t in (_paragraph_text(p).strip() for p in cell.iter() if _tag(p) == "p") if t) for cell in row if _tag(cell) == "tc"]
                for row in block.iter()
                if _tag(row) == "tr"
            ]
            parts.append(f"{table_title()}:\n{markdown_table(rows)}")

        elif tag == "sdt":
            # Its tables are not the body's: not numbered as docs.make_template counts.
            for content in block:
                if _tag(content) == "sdtContent":
                    parts += _blocks(content, lambda: "Таблица")

    return parts


def _paragraph_text(paragraph) -> str:
    """The text as it reads: a tab or a line break is a space, not two words run together."""
    pieces = []
    for node in paragraph.iter():
        tag = _tag(node)
        if tag == "t":
            pieces.append(node.text or "")
        elif tag in ("tab", "br", "cr"):
            pieces.append(" ")

    return "".join(pieces)


def _tag(node) -> str:
    return node.tag.rsplit("}", 1)[-1] if isinstance(node.tag, str) else ""


def _read_eml(path: Path) -> str:
    """A saved letter (a forwarded one is kept so): its headers, its text, its files by name."""
    from email import message_from_bytes, policy
    from html import unescape
    import re

    message = message_from_bytes(path.read_bytes(), policy=policy.default)
    lines = [f"{name}: {message[name]}" for name in ("From", "To", "Cc", "Date", "Subject") if message[name]]

    body = message.get_body(preferencelist=("plain", "html"))
    if body is not None:
        try:
            text = body.get_content()
        except (LookupError, UnicodeError):
            text = (body.get_payload(decode=True) or b"").decode("utf-8", errors="replace")
        if body.get_content_type() == "text/html":
            text = unescape(re.sub(r"<(script|style)\b.*?</\1>|<[^>]+>", " ", text, flags=re.S | re.I))
        lines += ["", text.strip()]

    files = [part.get_filename() for part in message.iter_attachments() if part.get_filename()]
    if files:
        lines += ["", "Вложения: " + ", ".join(files)]

    return "\n".join(lines)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []

    try:
        for sheet in workbook.worksheets:
            rows = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    rows.append([f"… ещё строки (показаны первые {MAX_SHEET_ROWS})"])
                    break

                cells = [cell_to_text(v) for v in row]
                if any(cells):
                    rows.append(cells)

            parts.append(f"## Лист «{sheet.title}»\n\n{markdown_table(rows)}")

    finally:
        workbook.close()

    return "\n\n".join(parts)


def _read_pdf(path: Path) -> str:
    import pdfplumber

    with pdfplumber.open(str(path)) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]

    text = "\n\n".join(p for p in pages if p.strip())
    if not text:
        raise ToolError("В PDF нет текстового слоя (похоже на скан)")

    return text


def _read_text(path: Path) -> str:
    raw = path.read_bytes()

    # Office exports on Russian Windows are often cp1251.
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)

        except UnicodeDecodeError:
            continue

    return raw.decode("utf-8", errors="replace")
