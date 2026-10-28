"""A sheet as people make them, read into a header and rows: shared by xlsx.read_table
and xlsx.aggregate so that both see the same table.

- A merged cell holds its value in every cell it covers (a region over three rows
  is the region of each of them), not in the first one only.
- Columns with the same header are «Сумма», «Сумма (2)»: a later one does not hide an
  earlier one.
- The header is the row the caller names (a title above the table is common).
- CSV: the separator is the one the file uses (`,` `;` tab `|`), the encoding UTF-8 or
  else cp1251 (what Russian Excel saves).
- A file that is not a table says so, not as a ZipFile error."""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path
from typing import Any

from common.tools import ToolError


# A workbook bigger than this is read without its merged cells: openpyxl keeps a whole
# workbook in memory to know them.
_FULL_READ_BYTES = 20 * 2**20


def read_sheet(path: Path, sheet: str | None = None, header_row: int = 1) -> tuple[list[str], list[tuple[Any, ...]]]:
    """The header (unique names) and the rows below it, empty rows left out."""
    suffix = path.suffix.lower()

    if suffix == ".xls":
        raise ToolError(f"«{path.name}» — старый формат Excel (.xls): сохраните его как .xlsx")

    if suffix in (".csv", ".tsv", ".txt"):
        values = _csv_rows(path)
    else:
        values = _xlsx_rows(path, sheet)

    values = values[header_row - 1 :]
    if not values:
        raise ToolError(f"В «{path.name}» нет строки {header_row}: лист пустой или короче")

    return unique_columns(values[0]), [row for row in values[1:] if any(v not in (None, "") for v in row)]


def unique_columns(header: tuple[Any, ...] | list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    columns = []

    for i, cell in enumerate(header):
        name = " ".join(str(cell).split()) if cell not in (None, "") else f"col_{i + 1}"
        seen[name] = seen.get(name, 0) + 1
        columns.append(name if seen[name] == 1 else f"{name} ({seen[name]})")

    return columns


def _xlsx_rows(path: Path, sheet: str | None) -> list[tuple[Any, ...]]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    full = path.stat().st_size <= _FULL_READ_BYTES

    try:
        workbook = load_workbook(str(path), read_only=not full, data_only=True)

    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as e:
        raise ToolError(f"«{path.name}» — не таблица Excel (.xlsx), прочитать её нельзя") from e

    # read_only keeps the file open until close(): on a Windows share an open file is locked.
    try:
        try:
            worksheet = workbook[sheet] if sheet else workbook.worksheets[0]

        except KeyError as e:
            raise ToolError(f"В «{path.name}» нет листа «{sheet}»; есть: {workbook.sheetnames}") from e

        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]

        if full:
            for merged in worksheet.merged_cells.ranges:
                value = worksheet.cell(merged.min_row, merged.min_col).value
                for r in range(merged.min_row, merged.max_row + 1):
                    for c in range(merged.min_col, merged.max_col + 1):
                        if r - 1 < len(rows) and c - 1 < len(rows[r - 1]):
                            rows[r - 1][c - 1] = value

        return [tuple(row) for row in rows]

    finally:
        workbook.close()


def _csv_rows(path: Path) -> list[tuple[Any, ...]]:
    raw = path.read_bytes()

    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break

        except UnicodeDecodeError:
            continue

    try:
        dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")

    except csv.Error:
        dialect = csv.excel

    return [tuple(row) for row in csv.reader(io.StringIO(text), dialect)]
