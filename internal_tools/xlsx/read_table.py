from __future__ import annotations

import asyncio
from datetime import date, datetime
from typing import Any

from pydantic import Field

from common.tools import ToolContext, ToolError, ToolInput, ToolOutput, tool


class ReadInput(ToolInput):
    path: str = Field(description="Файл .xlsx")
    sheet: str | None = Field(default=None, description="Лист; по умолчанию первый")
    header_row: int = Field(default=1, ge=1, description="Номер строки с заголовками")
    max_rows: int = Field(default=1000, ge=1, le=50000)


class ReadOutput(ToolOutput):
    columns: list[str]
    rows: list[dict[str, Any]]
    total_rows: int


@tool(
    name="xlsx.read_table",
    title="Прочитать таблицу Excel",
    card="Прочитать таблицу из «{path}»",
    description="Read an Excel sheet as a list of row objects keyed by the header row.",
    side_effect="read",
    capabilities=("fs.read:{path}",),
)
async def read_table(inp: ReadInput, ctx: ToolContext) -> ReadOutput:
    path = ctx.read_path(inp.path)
    if not path.is_file():
        raise ToolError(f"Файл `{inp.path}` не найден")

    return await asyncio.to_thread(_read, path, inp)


def _read(path, inp: ReadInput) -> ReadOutput:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)

    try:
        sheet = workbook[inp.sheet] if inp.sheet else workbook.worksheets[0]

    except KeyError as e:
        raise ToolError(f"Листа «{inp.sheet}» нет; есть: {workbook.sheetnames}") from e

    rows_iter = sheet.iter_rows(min_row=inp.header_row, values_only=True)
    header = next(rows_iter, None)
    if header is None:
        raise ToolError("Лист пустой")

    columns = [str(c).strip() if c is not None else f"col_{i + 1}" for i, c in enumerate(header)]
    rows: list[dict[str, Any]] = []
    total = 0

    for values in rows_iter:
        if all(v is None for v in values):
            continue

        total += 1
        if len(rows) < inp.max_rows:
            rows.append({columns[i]: _jsonable(v) for i, v in enumerate(values) if i < len(columns)})

    workbook.close()

    return ReadOutput(columns=columns, rows=rows, total_rows=total)


def _jsonable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return value
