"""Deterministic totals over spreadsheet rows: the LLM never adds numbers up."""

from __future__ import annotations

import asyncio
from collections import defaultdict
from typing import Any

from pydantic import BaseModel, Field

from common.tools import ToolContext, ToolError, ToolInput, ToolOutput, tool
from common.tools.refs import FileRefs, file_refs


class Input(ToolInput):
    inputs: FileRefs
    sum_columns: list[str] = Field(default_factory=list, description="Колонки, которые нужно просуммировать")
    group_by: str | None = Field(default=None, description="Колонка для группировки (например, Клиент)")
    top_n: int | None = Field(default=None, ge=1, description="Оставить N групп с наибольшей суммой первой из sum_columns")
    sheet: str | None = Field(default=None, description="Лист; по умолчанию первый")


class Group(BaseModel):
    key: str
    count: int
    sums: dict[str, float]


class Output(ToolOutput):
    files: int
    rows: int
    totals: dict[str, float]
    groups: list[Group]


@tool(
    name="xlsx.aggregate",
    title="Посчитать итоги по таблицам",
    card="Посчитать итоги: суммы {sum_columns}, группировка «{group_by}», топ {top_n}",
    description=(
        "Compute exact totals over the rows of one or more spreadsheets: row count, sums of columns, "
        "optional group-by with per-group counts and sums, optional top-N groups. Use this for any numbers "
        "in a report instead of asking the LLM to add them up."
    ),
    confirm=False,
)
async def aggregate(inp: Input, ctx: ToolContext) -> Output:
    paths = [ctx.read_path(ref) for ref in file_refs(inp.inputs)]
    if not paths:
        raise ToolError("Нет файлов для подсчёта")

    rows: list[dict[str, Any]] = []
    for path in paths:
        if not path.is_file():
            raise ToolError(f"Файл `{ctx.to_ref(path)}` не найден")

        rows += await asyncio.to_thread(_read_rows, path, inp.sheet)

    if rows:
        missing = [c for c in [*inp.sum_columns, *([inp.group_by] if inp.group_by else [])] if c not in rows[0]]
        if missing:
            raise ToolError(f"Нет колонок {missing}; есть: {list(rows[0])}")

    totals = {column: sum(_number(row.get(column)) for row in rows) for column in inp.sum_columns}

    groups: list[Group] = []
    if inp.group_by:
        buckets: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in rows:
            buckets[str(row.get(inp.group_by) or "—")].append(row)

        groups = [
            Group(key=key, count=len(items), sums={c: sum(_number(r.get(c)) for r in items) for c in inp.sum_columns})
            for key, items in buckets.items()
        ]

        if inp.sum_columns:
            groups.sort(key=lambda g: g.sums[inp.sum_columns[0]], reverse=True)

        if inp.top_n:
            groups = groups[: inp.top_n]

    return Output(files=len(paths), rows=len(rows), totals=totals, groups=groups)


def _read_rows(path, sheet: str | None) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        import csv

        with path.open(encoding="utf-8-sig", newline="") as fh:
            return list(csv.DictReader(fh))

    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.worksheets[0]

    values = worksheet.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else f"col_{i + 1}" for i, c in enumerate(next(values, []))]
    rows = [dict(zip(header, row)) for row in values if any(v is not None for v in row)]

    workbook.close()

    return rows


def _number(value: Any) -> float:
    if value is None or value == "":
        return 0.0

    if isinstance(value, (int, float)):
        return float(value)

    try:
        return float(str(value).replace(" ", "").replace(" ", "").replace(",", "."))

    except ValueError as e:
        raise ToolError(f"Значение {value!r} не число") from e
