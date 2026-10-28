"""Tables and data as people make them, as the stress test found them mishandled: merged
cells, a header twice, a title above the table, a totals row, a client written two ways,
a CSV of `;` in cp1251, numbers with thin spaces; a file that is not a table; outputs of
the same name; a model echoing a schema's field names; a template given null or a string
for its rows; a template replaced by accident; a notice in a dry run."""

from __future__ import annotations

import pytest

from common.tools import ToolContext, ToolError
from common.utils import parse_number
from executor.paths import WorkspacePaths
from internal_tools.xlsx.sheets import read_sheet
from tests.test_tools import _example, _template


def _workbook(path, rows, merge=None, sheet="Лист1"):
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    for cells in merge or []:
        worksheet.merge_cells(cells)
    workbook.save(str(path))

    return path


def test_a_merged_cell_is_the_value_of_every_row_it_covers(tmp_path):
    path = _workbook(tmp_path / "m.xlsx", [["Регион", "Сумма"], ["Север", 10], [None, 20], [None, 30], ["Юг", 5]], merge=["A2:A4"])

    columns, rows = read_sheet(path)

    assert columns == ["Регион", "Сумма"] and [row[0] for row in rows] == ["Север", "Север", "Север", "Юг"]


def test_a_header_twice_keeps_both_columns_and_a_title_above_is_skipped(tmp_path):
    path = _workbook(tmp_path / "d.xlsx", [["Продажи за март"], [], ["Клиент", "Сумма", "Сумма"], ["Альфа", 100, 1]])

    columns, rows = read_sheet(path, header_row=3)

    assert columns == ["Клиент", "Сумма", "Сумма (2)"] and rows == [("Альфа", 100, 1)]


@pytest.mark.parametrize(("text", "encoding"), [("Клиент;Сумма\nАльфа;1 200,50\nБета;300\n", "cp1251"), ("Клиент,Сумма\nАльфа,1200.5\n", "utf-8")])
def test_a_csv_is_read_with_its_own_separator_and_encoding(tmp_path, text, encoding):
    path = tmp_path / "t.csv"
    path.write_bytes(text.encode(encoding))

    columns, rows = read_sheet(path)

    assert columns == ["Клиент", "Сумма"] and rows[0][0] == "Альфа" and parse_number(rows[0][1]) == 1200.5


@pytest.mark.parametrize(("name", "content", "says"), [("fake.xlsx", b"just text", "не таблица Excel"), ("old.xls", b"\xd0\xcf", "старый формат")])
def test_a_file_that_is_not_a_table_says_so(tmp_path, name, content, says):
    path = tmp_path / name
    path.write_bytes(content)

    with pytest.raises(ToolError, match=says):
        read_sheet(path)


@pytest.mark.parametrize(
    ("value", "number"),
    [("1 000", 1000), ("1 000,50", 1000.5), ("1 234", 1234), ("1'234.5", 1234.5), ("1.234,56", 1234.56), ("1,234.56", 1234.56), ("12,5", 12.5)],
)
def test_numbers_as_tables_write_them(value, number):
    assert parse_number(value) == number


async def test_totals_leave_out_the_tables_own_total_row_and_join_one_client(settings, env):
    from internal_tools.xlsx.aggregate import Input, aggregate

    _workbook(env / "workspace" / "s.xlsx", [
        ["Клиент", "Сумма"], ["ООО Ромашка", 100], ["ООО  Ромашка ", 50], ["ооо ромашка", 1], [0, 7], ["Итого", 158],
    ])
    ctx = ToolContext(mode="chat", settings=settings, paths=WorkspacePaths(env / "workspace", settings.storage.data_dir))

    out = await aggregate(Input(inputs="s.xlsx", sum_columns=["Сумма"], group_by="Клиент"), ctx)

    assert out.totals == {"Сумма": 158} and out.total_rows_skipped == 1
    assert [(g.key, g.sums["Сумма"]) for g in out.groups] == [("ООО Ромашка", 151), ("0", 7)]


async def test_outputs_of_the_same_name_are_two_files_and_a_reserved_name_is_not_used(settings, env):
    ctx = ToolContext(mode="chat", settings=settings, paths=WorkspacePaths(env / "workspace", settings.storage.data_dir), session_id="ses_o")

    first = ctx.output_path("chart.png")
    first.write_bytes(b"1")
    second = ctx.output_path("chart.png")

    assert first.name == "chart.png" and second.name == "chart (2).png"
    assert ctx.output_path("CON", suffix=".xlsx").name == "_CON.xlsx"


async def test_a_sheet_name_excel_refuses_is_made_one_it_takes(settings, env):
    from openpyxl import load_workbook

    from internal_tools.xlsx.write_table import WriteInput, write_table

    ctx = ToolContext(mode="chat", settings=settings, paths=WorkspacePaths(env / "workspace", settings.storage.data_dir), session_id="ses_w")

    out = await write_table(WriteInput(rows=[{"a": 1}], sheet="Итоги 10/2026"), ctx)

    assert load_workbook(str(ctx.read_path(out.path))).sheetnames == ["Итоги 10_2026"]


async def test_a_model_echoing_the_schemas_field_names_is_not_an_answer(monkeypatch, settings, env):
    from internal_tools.llm import extract as extract_tool

    seen = []

    class Model:
        async def structured(self, messages, *, schema, name, validate):
            for answer in ({"services": [{"qty": "price", "amount": "name"}], "total": "name"}, {"services": [{"qty": "1", "amount": "40000"}], "total": "47500"}):
                try:
                    seen.append(validate(answer))
                    return answer
                except ValueError as e:
                    seen.append(str(e))

    monkeypatch.setattr(extract_tool, "instruct", lambda ctx: Model())
    schema = {"type": "object", "properties": {
        "services": {"type": "array", "items": {"type": "object", "properties": {"name": {"type": "string"}, "qty": {"type": "string"}, "price": {"type": "string"}, "amount": {"type": "string"}}}},
        "total": {"type": "string"},
    }}

    out = await extract_tool.extract(extract_tool.ExtractInput(inputs="Аудит сети, 1 шт, 40 000; итого 47 500", instruction="услуги", output_schema=schema),
        ToolContext(mode="chat", settings=settings, paths=WorkspacePaths(env / "workspace", settings.storage.data_dir)),
    )

    assert "имена полей схемы" in seen[0] and out.data["total"] == "47500"


@pytest.mark.parametrize(
    ("values", "says"),
    [
        ({"total": None, "items": []}, "пустое \\(null\\)"),
        ({"total": ["a", "b"], "items": []}, "список"),
        ({"total": 1, "items": "name"}, "нужен список"),
    ],
)
def test_a_template_given_null_a_list_or_a_string_for_rows_fails(tmp_path, values, says):
    from internal_tools.docs.fill_template import _render

    with pytest.raises(ToolError, match=says):
        _render(_template(tmp_path / "t.docx"), values, {}, tmp_path / "out.docx")


async def test_a_template_is_replaced_only_when_asked(settings, env):
    from internal_tools.docs.make_template import Input, TemplateField, make_template

    _example(env / "workspace" / "example.docx")
    ctx = ToolContext(mode="chat", settings=settings, paths=WorkspacePaths(env / "workspace", settings.storage.data_dir), session_id="ses_t")
    markup = {"example": "example.docx", "name": "weekly", "fields": [TemplateField(name="report_date", description="Дата", example="7 декабря 2026 г.")]}

    await make_template(Input(**markup), ctx)
    with pytest.raises(ToolError, match="уже есть"):
        await make_template(Input(**markup), ctx)

    assert (await make_template(Input(**markup, overwrite=True), ctx)).template == "data:templates/weekly.docx"


async def test_a_notice_in_a_dry_run_is_not_sent():
    from internal_tools.user.notify import NotifyInput, notify

    ctx = ToolContext(mode="run", settings=None, paths=None, dry_run=True)
    out = await notify(NotifyInput(message="Готово"), ctx)

    assert out.sent is False and ctx.notifications == ["Готово"] and "не отправлено" in ctx.logs[0]
