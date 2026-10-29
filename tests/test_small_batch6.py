"""The small things the stress test found: a server's answer as text, a folder written with
a slash on a server of dots, a letter without a Date, a forwarded .eml and a docx footer
read, cards without «— —», an absolute path hidden in a list."""

from __future__ import annotations

import pytest

from compiler.validate import SpecError, validate_spec
from internal_tools.mail.adapters import folders, imap_smtp
from tests.test_mail_folders import FakeFolders


def test_a_servers_answer_reads_as_text():
    assert folders.said([b"EXAMINE failed. No such mailbox"]) == "EXAMINE failed. No such mailbox"
    assert folders.said([None]) == "сервер не объяснил"


@pytest.mark.parametrize(
    ("rows", "name", "on_server"),
    [
        (['() "." "INBOX"', '() "." "Рассылки"'], "Рассылки/Новости", "Рассылки.Новости"),   # a server of dots
        (['() "/" "INBOX"'], "Рассылки/Новости", "Рассылки/Новости"),                        # of slashes: as it is
        (['() "." "INBOX"', '() "." "Итоги 10/2026"'], "Итоги 10/2026", "Итоги 10/2026"),     # a folder of that very name
    ],
)
def test_a_folder_written_with_a_slash_is_the_servers_nested_one(rows, name, on_server):
    assert folders.on_server(FakeFolders(*rows), name) == on_server


def test_a_letter_without_a_date_has_the_one_the_server_got_it():
    info = b'1 (UID 5 FLAGS () INTERNALDATE "28-Oct-2026 09:15:00 +0300" BODY[HEADER.FIELDS (FROM)] {20}'

    item = imap_smtp._header_item(5, b"From: a@b.ru\r\n\r\n", [], imap_smtp._received(info))

    assert item["date"].startswith("2026-10-28")


def test_a_saved_letter_is_read(tmp_path):
    from email.message import EmailMessage

    from internal_tools.docs.reader import read_document

    letter = EmailMessage()
    letter["From"], letter["Subject"] = "anna@umail.com", "Акт сверки"
    letter.set_content("Акт во вложении.")
    letter.add_attachment(b"1", maintype="text", subtype="plain", filename="акт.txt")
    path = tmp_path / "Акт.eml"
    path.write_bytes(letter.as_bytes())

    text = read_document(path)

    assert "Subject: Акт сверки" in text and "Акт во вложении." in text and "Вложения: акт.txt" in text


def test_a_docx_header_and_footer_are_read(tmp_path):
    from docx import Document

    from internal_tools.docs.reader import read_document

    document = Document()
    document.sections[0].header.paragraphs[0].text = "ООО «Глобус» — бухгалтерия"
    document.sections[0].footer.paragraphs[0].text = "Подпись: ______"
    document.add_paragraph("Акт № 5")
    path = tmp_path / "a.docx"
    document.save(str(path))

    assert read_document(path).split("\n\n") == ["Верхний колонтитул: ООО «Глобус» — бухгалтерия", "Акт № 5", "Нижний колонтитул: Подпись: ______"]


def test_a_card_leaves_out_what_is_not_given_and_shows_defaults():
    from common.tools import discover
    from compiler.cards import _format_card

    manifests = {name: spec.manifest() for name, spec in discover().items()}
    mark, forward = manifests["mail.mark"], manifests["mail.forward"]

    assert _format_card(mark["card"], {"account": "Test", "uids": [3]}, fallback="x", schema=mark["input_schema"]) == "Пометить письма 3 в «Test»"
    assert _format_card(forward["card"], {"uid": 7, "to": ["a@b.ru"]}, fallback="x", schema=forward["input_schema"]) == "Переслать письмо 7 адресатам a@b.ru, из папки «INBOX»"


def test_an_absolute_path_anywhere_in_the_arguments_is_refused():
    from common.tools import discover

    catalog = {name: spec.manifest() for name, spec in discover().items()}
    send = {"account": "Test", "to": ["a@umail.com"], "subject": "s", "body": "b", "attachments": ["Отчёты/a.txt", "/etc/passwd"]}

    with pytest.raises(SpecError, match="абсолютный путь '/etc/passwd'"):
        validate_spec({"title": "t", "steps": [{"id": "s", "tool": "mail.send", "args": send}]}, catalog)

    # Text that starts with a slash is not a path.
    validate_spec({"title": "t", "steps": [{"id": "n", "tool": "user.notify", "args": {"message": "/start — команда бота"}}]}, catalog)


def test_a_docx_reads_tabs_form_fields_and_header_tables(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    from internal_tools.docs.reader import read_document

    document = Document()
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.header.add_table(rows=1, cols=2, width=section.page_width).rows[0].cells[0].text = "ИНН 7700000001"
    section.even_page_header.paragraphs[0].text = "Чётная страница"

    paragraph = document.add_paragraph()
    paragraph.add_run("Итого:")
    paragraph.add_run().add_tab()
    paragraph.add_run("150")
    cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
    cell.paragraphs[0].add_run("Ром")
    cell.paragraphs[0].add_run("ашка")
    document.element.body.append(parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>Подписант: Иванов</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    ))
    path = tmp_path / "a.docx"
    document.save(str(path))

    text = read_document(path)

    assert "Итого: 150" in text and "| Ромашка |" in text and "Подписант: Иванов" in text
    assert "ИНН 7700000001" in text and "Верхний колонтитул: Чётная страница" in text
