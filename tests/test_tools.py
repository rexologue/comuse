"""Deterministic tools that frozen tasks rely on instead of the LLM."""

from __future__ import annotations

import pytest

from common.tools import ToolError, discover
from internal_tools.data.map_rows import Input as MapInput, map_rows


GROUPS = [
    {"key": "АО Техноресурс", "count": 5, "sums": {"Сумма": 334500.0}},
    {"key": "ООО Ромашка", "count": 4, "sums": {"Сумма": 289000.0}},
    {"key": "ИП Смирнов", "count": 4, "sums": {"Сумма": 229000.0}},
]


async def test_map_rows_renames_formats_and_limits():
    out = await map_rows(
        MapInput(rows=GROUPS, fields={"client": "row.key", "amount": "row.sums['Сумма'] | format_number"}, limit=2),
        ctx=None,
    )

    assert out.rows == [
        {"client": "АО Техноресурс", "amount": "334 500"},
        {"client": "ООО Ромашка", "amount": "289 000"},
    ]


async def test_map_rows_reports_bad_expression():
    with pytest.raises(ToolError):
        await map_rows(MapInput(rows=GROUPS, fields={"x": "row.nope.deeper"}), ctx=None)


def test_catalog_has_the_deterministic_tools():
    names = set(discover())

    assert {"xlsx.aggregate", "data.map_rows", "docs.fill_template"} <= names


def _template(path):
    from docx import Document

    document = Document()
    document.add_paragraph("Итого: {{ total }}")
    table = document.add_table(rows=4, cols=1)
    table.rows[0].cells[0].text = "{%tr for row in items %}"
    table.rows[1].cells[0].text = "{{ row.name }}"
    table.rows[2].cells[0].text = "{%tr endfor %}"
    table.rows[3].cells[0].text = "конец"
    document.save(str(path))

    return path


def test_fill_template_checks_fields_and_row_keys(tmp_path):
    from docx import Document

    from internal_tools.docs.fill_template import _render

    template = _template(tmp_path / "t.docx")
    out = tmp_path / "out.docx"

    with pytest.raises(ToolError, match="total"):
        _render(template, {"items": []}, {}, out)

    with pytest.raises(ToolError, match="data.map_rows"):
        _render(template, {"total": 1, "items": [{"key": "a"}]}, {}, out)

    _render(template, {"total": "3", "items": [{"name": "Ромашка"}]}, {}, out)
    text = [p.text for p in Document(str(out)).paragraphs] + [c.text for t in Document(str(out)).tables for r in t.rows for c in r.cells]

    assert "Итого: 3" in text and "Ромашка" in text


def _example(path):
    """A filled-in report: a date in the text, a sales table with one header row."""
    from docx import Document

    document = Document()
    document.add_paragraph("Отчёт за 7 декабря 2026 г.")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Клиент", "Сумма"), ("Ромашка", "100"), ("Вектор", "250")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(str(path))

    return path


def test_make_template_applies_the_markup_and_fills_back(tmp_path):
    from docx import Document

    from internal_tools.docs.fill_template import _render
    from internal_tools.docs.make_template import TemplateField, TemplateTable, _make

    template = tmp_path / "t.docx"
    result = _make(
        _example(tmp_path / "example.docx"),
        template,
        [TemplateField(name="report_date", description="Дата отчёта", example="7 декабря 2026 г.")],
        [TemplateTable(table=1, name="sales", header_rows=1, columns=["client", "amount"])],
    )

    assert result["example_context"] == {
        "report_date": "7 декабря 2026 г.",
        "sales": [{"client": "Ромашка", "amount": "100"}, {"client": "Вектор", "amount": "250"}],
    }

    out = tmp_path / "out.docx"
    _render(template, {"report_date": "8 декабря 2026 г.", "sales": [{"client": "Альфа", "amount": "5"}]}, {}, out)
    document = Document(str(out))
    cells = [[c.text for c in r.cells] for r in document.tables[0].rows]

    assert document.paragraphs[0].text == "Отчёт за 8 декабря 2026 г."
    assert cells == [["Клиент", "Сумма"], ["Альфа", "5"]]


def test_make_template_names_every_problem_of_the_markup(tmp_path):
    from internal_tools.docs.make_template import TemplateField, TemplateTable, _make

    target = tmp_path / "t.docx"

    with pytest.raises(ToolError) as e:
        _make(
            _example(tmp_path / "example.docx"),
            target,
            [TemplateField(name="Date", description="", example="9 мая")],
            [
                TemplateTable(table=1, name="sales", header_rows=1, columns=["client", "amount", "vat"]),
                TemplateTable(table=3, name="other", header_rows=0, columns=["x"]),
            ],
        )

    message = str(e.value)
    assert "«Date»" in message and "«9 мая»" in message
    assert "колонок 2, а ключей 3" in message and "T3 нет" in message
    assert not target.exists()


def test_docx_tables_are_numbered_as_make_template_counts_them(tmp_path):
    from common.documents import read_document

    text = read_document(_example(tmp_path / "example.docx"))

    assert "Таблица T1:\n| Клиент | Сумма |" in text


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("a/b.xlsx", ["a/b.xlsx"]),
        ('["a/b.xlsx", "c/d.xlsx"]', ["a/b.xlsx", "c/d.xlsx"]),
        ("a/b.xlsx, c/d.xlsx", ["a/b.xlsx", "c/d.xlsx"]),
        ({"files": [{"path": "a.xlsx", "name": "a.xlsx"}]}, ["a.xlsx"]),
        (["a.xlsx", {"path": "b.xlsx"}], ["a.xlsx", "b.xlsx"]),
        ("Отчёты/Иванов, Петров.docx", ["Отчёты/Иванов, Петров.docx"]),
    ],
)
def test_file_refs_accepts_every_shape(value, expected):
    from common.tools.refs import file_refs

    assert file_refs(value) == expected


def test_aggregate_inputs_have_a_schema():
    schema = discover()["xlsx.aggregate"].manifest()["input_schema"]["properties"]["inputs"]

    assert "anyOf" in schema


async def test_move_in_a_dry_run_only_says_what_it_would_do(settings):
    from common import tunables
    from common.tools import ToolContext
    from executor.paths import WorkspacePaths
    from internal_tools.fs.transfer import Input, move_file

    workspace = tunables.get("workspace_dir")
    (workspace / "Входящие").mkdir()
    (workspace / "Входящие" / "a.pdf").write_bytes(b"%PDF")

    ctx = ToolContext(
        mode="run", settings=settings, paths=WorkspacePaths(workspace, settings.storage.data_dir), llm=None, dry_run=True,
    )
    out = await move_file(Input(src="Входящие/a.pdf", dst="Счета/"), ctx)

    assert out.path == "Счета/a.pdf"
    assert (workspace / "Входящие" / "a.pdf").exists() and not (workspace / "Счета").exists()
    assert ctx.logs == ["Пробный запуск: переложил бы «Входящие/a.pdf» в «Счета/a.pdf»"]


async def test_fs_list_stays_in_its_folder(settings, tmp_path):
    from common import tunables
    from common.tools import ToolContext
    from executor.paths import WorkspacePaths
    from internal_tools.fs.list import Input, list_files

    workspace = tunables.get("workspace_dir")
    (workspace / "a.txt").write_text("a")
    (tmp_path / "outside.txt").write_text("secret")
    (workspace / "link.txt").symlink_to(tmp_path / "outside.txt")
    (settings.storage.data_dir / "outputs").mkdir(parents=True, exist_ok=True)

    ctx = ToolContext(mode="chat", settings=settings, paths=WorkspacePaths(workspace, settings.storage.data_dir), llm=None)

    for pattern in ("../*", "/etc/*", "x/../../*"):
        with pytest.raises(ToolError, match="без `..`"):
            await list_files(Input(folder="data:outputs", pattern=pattern), ctx)

    # A link out of the workspace is left out, the rest is listed.
    listed = await list_files(Input(), ctx)
    assert [f.path for f in listed.files] == ["a.txt"]
